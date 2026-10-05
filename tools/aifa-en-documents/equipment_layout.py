from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def improve_equipment_layout(path):
    doc=Document(path)
    weights={2:[1.2,.55,.5,.7,.6,.6,.8,1.4],3:[1.2,.55,.5,.7,.6,.6,.8,1.4],6:[1.3,1.1,.65,2.55,1.05],7:[1.3,1.1,.65,2.55,1.05],8:[1.7,.65,3],9:[1.3,.8,.8,1,1,1.6,1.2,.7],10:[.7,1.6,1.7,1.8]}
    for i,table in enumerate(doc.tables):
        table.autofit=False
        total=sum(c.width for c in table.columns)
        ratio=weights.get(i,[1]*len(table.columns))
        widths=[int(total*w/sum(ratio)) for w in ratio]
        for col,width in zip(table.columns,widths):col.width=width
        for row in table.rows:
            for cell,width in zip(row.cells,widths):
                cell.width=width
                pr=cell._tc.get_or_add_tcPr()
                borders=pr.find(qn('w:tcBorders'))
                if borders is None:borders=OxmlElement('w:tcBorders');pr.append(borders)
                for edge in ['top','left','bottom','right']:
                    element=borders.find(qn('w:'+edge))
                    if element is None:element=OxmlElement('w:'+edge);borders.append(element)
                    for key,value in [('val','single'),('sz','4'),('color','000000')]:element.set(qn('w:'+key),value)
                if i==9:
                    for p in cell.paragraphs:
                        for run in p.runs:run.font.size=Pt(9)
        pr=table.rows[0]._tr.get_or_add_trPr()
        if pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
    doc.paragraphs[59].paragraph_format.keep_together=True
    doc.core_properties.title='AI Fantasy Adventure – Equipment Catalogue v1.0'
    doc.save(path)
